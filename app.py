import streamlit as st
st.set_page_config(layout="wide")  # Must be the first Streamlit command

import pdfplumber  # For PDF extraction
import os
from dotenv import load_dotenv
from langchain_groq import ChatGroq
from docx import Document  # For Word document extraction
import base64
from datetime import datetime
import json
from fpdf import FPDF
from io import BytesIO

# Azure Computer Vision imports
from azure.cognitiveservices.vision.computervision import ComputerVisionClient
from msrest.authentication import CognitiveServicesCredentials

# Load environment variables
load_dotenv()
API_KEY = os.getenv("GROQ_API_KEY")
llm = ChatGroq(groq_api_key=API_KEY)

# Load Azure credentials
AZURE_ENDPOINT = os.getenv("AZURE_ENDPOINT")
AZURE_API_KEY = os.getenv("AZURE_API_KEY")

# Azure OCR extraction from image bytes
def extract_text_from_image(image_bytes):
    if not AZURE_ENDPOINT or not AZURE_API_KEY:
        return "Error: Azure credentials not set."
    
    client = ComputerVisionClient(AZURE_ENDPOINT, CognitiveServicesCredentials(AZURE_API_KEY))
    try:
        img_stream = BytesIO(image_bytes)
        analysis = client.read_in_stream(img_stream, raw=True)
        operation_location = analysis.headers.get("Operation-Location")
        if not operation_location:
            return "Error: No operation location returned."
        operation_id = operation_location.split("/")[-1]

        # Poll for the result
        while True:
            result = client.get_read_result(operation_id)
            if result.status not in ["notStarted", "running"]:
                break

        if result.status == "succeeded":
            extracted_text = "\n".join([line.text for page in result.analyze_result.read_results for line in page.lines])
            return extracted_text
        else:
            return "Error: OCR extraction did not succeed."
    except Exception as e:
        return f"Error during OCR extraction: {str(e)}"

# Extract text from PDF
def extract_text_from_pdf(pdf_file):
    text = ""
    try:
        with pdfplumber.open(pdf_file) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text
                else:
                    page_image = page.to_image()
                    img_bytes = page_image.to_bytes(format="PNG")
                    text += extract_text_from_image(img_bytes)
    except Exception as e:
        text = f"Error reading PDF: {str(e)}"
    return text

# Extract text from Word documents
def extract_text_from_docx(docx_file):
    try:
        doc = Document(docx_file)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        return f"Error reading Word document: {str(e)}"

# Call to GroqCloud API with refined template
def call_groqcloud_api(summary):
    try:
        prompt = f"""
        Extract key patient information from the document.
        Return a JSON with fields:
        - Name
        - Age (convert from Birthdate if available using today's date)
        - Birthdate
        - Address (parse into Street, City, State, Pincode)
        - Contact Information (Phone, Email)
        - Medical History
        - Diagnosis
        - Prescription Details
        - Any other relevant patient-specific information
        - Automatically calculate age from the birthdate if available (use today's date as reference)
        - Parse addresses into structured fields
        - Add custom fields based on content (e.g., emergency contacts, insurance details)
        - Translate non-English content to English (except Name)
        [IMPORTANT] Return only JSON. No explanations or notes.
        
        Document:
        {summary}
        """
        response = llm.invoke(prompt)
        return response.content
    except Exception as e:
        return f"Error: {str(e)}"

# Generate PDF from JSON response
def generate_pdf_from_json(json_response, file_name="extracted_data.pdf"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Patient Information Report", ln=True, align='C')
    pdf.ln(10)

    try:
        data_dict = json.loads(json_response)
    except Exception as e:
        data_dict = {"Error": "Invalid JSON data"}

    pdf.set_font("Arial", "B", 12)
    pdf.cell(50, 10, "Field", border=1)
    pdf.cell(0, 10, "Value", border=1, ln=True)

    pdf.set_font("Arial", "", 12)
    for key, value in data_dict.items():
        pdf.cell(50, 10, str(key), border=1)
        pdf.cell(0, 10, str(value), border=1, ln=True)

    pdf.ln(10)
    diagnosis_summary = data_dict.get("Diagnosis", "No diagnosis available.")
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, "Diagnosis Summary:", ln=True)
    pdf.set_font("Arial", "", 12)
    pdf.multi_cell(0, 10, str(diagnosis_summary))

    pdf.output(file_name)

# Display PDF in Streamlit
def display_pdf(file_path):
    with open(file_path, "rb") as f:
        base64_pdf = base64.b64encode(f.read()).decode("utf-8")
    pdf_display = f'<iframe src="data:application/pdf;base64,{base64_pdf}" width="100%" height="500px" type="application/pdf"></iframe>'
    st.markdown(pdf_display, unsafe_allow_html=True)

# Main app logic
def main():
    st.title("Patient Document Parser & AI Extractor")

    uploaded_file = st.file_uploader("Upload Document", type=["pdf", "png", "jpg", "jpeg", "docx"])
    col1, col2 = st.columns([1, 1])

    extracted_text = ""
    api_response_text = ""

    if uploaded_file:
        file_type = uploaded_file.type
        with col1:
            if "pdf" in file_type:
                st.write("### Uploaded PDF:")
                temp_file_path = f"temp_{uploaded_file.name}"
                with open(temp_file_path, "wb") as f:
                    f.write(uploaded_file.read())
                display_pdf(temp_file_path)
                extracted_text = extract_text_from_pdf(temp_file_path)

            elif "docx" in file_type:
                st.write("### Uploaded Word Document:")
                extracted_text = extract_text_from_docx(uploaded_file)
                st.text_area("Preview", extracted_text, height=300)

            else:
                image_bytes = uploaded_file.getvalue()
                extracted_text = extract_text_from_image(image_bytes)
                st.image(uploaded_file, caption="Uploaded Image", use_column_width=True)

        # Summarize extracted text
        if extracted_text:
            summary_prompt = "Summarize this text while keeping all key details:\n" + extracted_text
            summary = llm.invoke(summary_prompt)

            # Call API with summarized text
            api_response_text = call_groqcloud_api(summary)

    with col2:
        st.write("### Extracted AI Response (JSON):")
        st.text_area("", api_response_text, height=400)

        if api_response_text:
            download_format = st.selectbox("Download Format", ["JSON", "Text", "PDF"], index=0)
            if download_format == "JSON":
                st.download_button("Download JSON", api_response_text, file_name="extracted_data.json", mime="application/json")
            elif download_format == "Text":
                st.download_button("Download Text", extracted_text, file_name="extracted_data.txt", mime="text/plain")
            elif download_format == "PDF":
                pdf_file_name = "extracted_data.pdf"
                generate_pdf_from_json(api_response_text, pdf_file_name)
                with open(pdf_file_name, "rb") as f:
                    st.download_button("Download PDF", f.read(), file_name=pdf_file_name, mime="application/pdf")

if __name__ == "__main__":
    main()
